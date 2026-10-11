"""Generate fictional fixtures in all three supported formats."""
from pathlib import Path
import os

from docx import Document
from openpyxl import Workbook
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "frontend" / "public" / "demo"
OUT.mkdir(parents=True, exist_ok=True)

TENDER = [
    "УЧЕБНЫЙ КОМПЛЕКТ DEMO-001",
    "Закупка ноутбуков. Все компании и условия в этом комплекте вымышлены.",
    "Оперативная память — не менее 16 ГБ.",
    "Накопитель — не менее 512 ГБ.",
    "Гарантия — не менее 24 месяцев.",
    "Документы — сертификат соответствия и авторизационное письмо.",
    "Срок поставки — не более 10 дней.",
    "Сервисный центр — на территории Узбекистана.",
]
BIDS = [
    ["УЧЕБНАЯ ЗАЯВКА DEMO-001 — ООО «Альфа»", "Все сведения вымышлены.",
     "Оперативная память: 16 ГБ.", "Накопитель: 512 ГБ.", "Гарантия: 24 месяца.",
     "Документы: сертификат соответствия и авторизационное письмо приложены.",
     "Срок поставки: 7 дней.", "Сервисный центр: Ташкент, Узбекистан."],
    ["УЧЕБНАЯ ЗАЯВКА DEMO-001 — ООО «Бета»", "Все сведения вымышлены.",
     "RAM: 8 GB.", "Storage: 512 GB.", "Warranty: 12 months.",
     "Документы: приложен только сертификат соответствия.", "Delivery: 14 days."],
    ["УЧЕБНАЯ ЗАЯВКА DEMO-001 — ООО «Гамма»", "Все сведения вымышлены.",
     "Оперативная память: 32 ГБ.", "Накопитель: 1024 ГБ.", "Гарантия: 24 месяца.",
     "Документы: сертификат соответствия и авторизационное письмо приложены.",
     "Срок поставки: по согласованию.", "Сервисный центр: отсутствует."],
]


def docx(name, lines):
    document = Document()
    for index, line in enumerate(lines):
        document.add_paragraph(line, style="Title" if index == 0 else None)
    document.save(OUT / name)


docx("tender.docx", TENDER)
docx("bid-a.docx", BIDS[0])
font_paths = [
    Path(os.getenv("DEMO_FONT", "C:/Windows/Fonts/arial.ttf")),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
]
font = next((path for path in font_paths if path.exists()), None)
if not font:
    raise RuntimeError("Provide a Unicode TTF file via DEMO_FONT")
pdfmetrics.registerFont(TTFont("DemoUnicode", str(font)))
pdf = canvas.Canvas(str(OUT / "bid-b.pdf"))
pdf.setFont("DemoUnicode", 11)
for index, line in enumerate(BIDS[1]):
    pdf.drawString(45, 780 - index * 32, line)
pdf.save()
workbook = Workbook()
workbook.active.title = "Заявка"
for line in BIDS[2]:
    workbook.active.append([line])
workbook.active.column_dimensions["A"].width = 95
workbook.save(OUT / "bid-c.xlsx")
print("Generated tender.docx, bid-a.docx, bid-b.pdf, bid-c.xlsx")
