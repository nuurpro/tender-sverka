from pathlib import Path

import pdfplumber
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

MAX_PACKAGE_CHARS = 100_000
SUPPORTED = {".pdf", ".docx", ".xlsx"}


def parse_package(files: list[dict]) -> list[dict]:
    blocks: list[dict] = []
    total = 0
    for file in files:
        path = Path(file["path"])
        file_blocks: list[dict] = []

        def add(suffix: str, location: str, text: str, page: int | None = None):
            nonlocal total
            text = text.strip()
            if not text:
                return
            total += len(text)
            if total > MAX_PACKAGE_CHARS:
                raise ValueError("Пакет превышает 100 000 символов извлечённого текста")
            file_blocks.append({
                "id": f"{file['id']}:{suffix}", "filename": file["name"],
                "location": location, "page": page, "text": text,
            })

        try:
            if path.suffix.lower() == ".pdf":
                with pdfplumber.open(path) as pdf:
                    for page in pdf.pages:
                        text = page.extract_text() or ""
                        if not text.strip():
                            raise ValueError(f"{file['name']}, страница {page.page_number}: нет текста; сканы не поддерживаются")
                        add(f"p{page.page_number}", f"Страница {page.page_number}", text, page.page_number)
            elif path.suffix.lower() == ".docx":
                document = Document(path)
                paragraph_number = table_number = 0
                for element in document.iter_inner_content():
                    if isinstance(element, Paragraph):
                        paragraph_number += 1
                        add(f"para{paragraph_number}", f"Абзац {paragraph_number}", element.text)
                    elif isinstance(element, Table):
                        table_number += 1
                        for row_number, row in enumerate(element.rows, 1):
                            text = " | ".join(cell.text for cell in row.cells)
                            add(f"t{table_number}:r{row_number}", f"Таблица {table_number}, строка {row_number}", text)
            elif path.suffix.lower() == ".xlsx":
                workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
                try:
                    for sheet_number, sheet in enumerate(workbook, 1):
                        for row in sheet.iter_rows():
                            values = []
                            for cell in row:
                                if cell.value is None:
                                    continue
                                value = "[Формула не обработана]" if cell.data_type == "f" else str(cell.value)
                                values.append(f"{cell.coordinate}: {value}")
                            if values:
                                row_number = row[0].row
                                add(f"s{sheet_number}:r{row_number}", f"Лист «{sheet.title}», строка {row_number}", " | ".join(values))
                finally:
                    workbook.close()
            else:
                raise ValueError("Поддерживаются только текстовые PDF, DOCX и XLSX")
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Не удалось прочитать «{file['name']}». Проверьте формат и защиту файла.") from exc
        if not file_blocks:
            raise ValueError(f"В «{file['name']}» нет извлекаемого текста")
        blocks.extend(file_blocks)
    if not blocks:
        raise ValueError("Пакет документов пуст")
    return blocks
