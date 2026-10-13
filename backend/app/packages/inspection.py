"""Measured inspection only. Extraction/OCR/tokenizer stages remain explicit gates."""
import hashlib
from importlib.metadata import version
from io import BytesIO
import logging
import json
import os
from pathlib import Path
import subprocess
import sys
from uuid import uuid4

from sqlalchemy import select

from app.workspace.database import tenant_transaction
from app.workspace.models import Document, Job, PackageInspection, Project
from app.workspace.repository import Repository, ValidationError, canonical_json
from .service import PackageService,nested_package
from .storage import StorageUnavailable

POLICY_VERSION='inspection-v1'


def parser_identity():
    return {'policy':POLICY_VERSION,'pdfplumber':version('pdfplumber'),'python-docx':version('python-docx'),
            'openpyxl':version('openpyxl'),'options':{'text_probe':True},'model':'none','tokenizer':'pending'}


def inspection_profile(settings,language,*,readiness=None):
    from app.documents.runner import DocumentRunnerError,ensure_runner_available
    from app.documents.service import processing_profile
    native=None
    try:
        native=processing_profile(settings,language,probe=(readiness or ensure_runner_available)(settings)).model_dump(mode='json')
    except DocumentRunnerError:
        pass  # Original upload does not depend on a running parser.
    return {'policy':'inspection-v3-ocr' if native and native['version']=='image-ocr-v2' else 'inspection-v2-native','language':language,'model':'none','tokenizer':'pending',
        'native_profile':native,'runner_status':'available' if native else 'unavailable',
        'options':{'text_probe':True,'max_pages':settings.max_package_pages,'max_tokens':settings.max_package_tokens}}


def inspect_document(data,mime,max_pages=1000,max_tokens=250_000):
    report={'page_count':None,'token_count':None,'text_probe_bytes':None,'token_upper_bound':None,
            'coverage':'pending','blockers':[],'warnings':[]}
    text=None
    try:
        if mime=='application/pdf':
            import pdfplumber
            with pdfplumber.open(BytesIO(data)) as pdf:
                report['page_count']=len(pdf.pages)
                if report['page_count']>max_pages:
                    report['blockers'].append('page_limit_exceeded')
                    return report
                text='\n'.join(page.extract_text() or '' for page in pdf.pages)
                if not text.strip():
                    report['coverage']='ocr_pending';report['blockers'].append('ocr_pending')
                else:
                    report['coverage']='probe'
        elif mime.endswith('wordprocessingml.document'):
            from docx import Document as WordDocument
            doc=WordDocument(BytesIO(data))
            text='\n'.join(p.text for p in doc.paragraphs)
            report['coverage']='probe';report['warnings'].append('office_layout_and_full_coverage_pending')
        elif mime.endswith('spreadsheetml.sheet'):
            from openpyxl import load_workbook
            workbook=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
            try:
                # Formula source is counted as inert text; formulas are never run.
                text='\n'.join(str(cell) for sheet in workbook for row in sheet.iter_rows(values_only=True) for cell in row if cell is not None)
            finally:
                workbook.close()
            report['coverage']='probe';report['warnings'].append('office_layout_and_full_coverage_pending')
        elif mime=='text/plain':
            text=data.decode('utf-8-sig');report['coverage']='probe'
        else:
            report['blockers'].append('adapter_unavailable')
        if text is not None:
            report['text_probe_bytes']=len(text.encode('utf-8'))
            # This is a provisional conservative byte-based bound of the probe,
            # never an accurate model token count or whole-extraction acceptance.
            report['token_upper_bound']=report['text_probe_bytes']
            if report['token_upper_bound']>max_tokens:
                report['blockers'].append('token_probe_bound_exceeded')
            report['warnings'].append('token_bound_provisional_utf8_probe_bytes')
    except Exception:
        report['coverage']='failed';report['blockers'].append('parser_failed')
    if report['page_count'] is None:
        report['blockers'].append('page_count_pending')
    report['blockers'].extend(['extraction_pending','tokenizer_profile_pending'])
    return report


def inspection_handler(settings,*,packages=None,parser=None,readiness=None):
    if settings.environment!='test' and (parser is not None or readiness is not None):
        raise ValueError('Inspection adapter injection requires the explicit test environment')
    from app.documents.runner import run_document,ensure_runner_available,DocumentRunnerError
    from app.documents.schemas import parse_profile
    from app.documents.service import processing_profile
    from .service import authorized
    parser=parser or run_document
    readiness=readiness or ensure_runner_available
    def handle(context):
        service=context.service;org=service.organization_id
        package_id=context.payload.get('package_id');actor=context.payload.get('actor_id')
        pipeline=packages or PackageService(service.engine,settings)
        with tenant_transaction(service.engine,org) as db:
            job=db.get(Job,context.job_id);project_id=job.project_id
        frozen=context.payload.get('inspection_profile')
        if not isinstance(frozen,dict) or not frozen.get('native_profile'):
            raise ValidationError('Frozen inspection has no available isolated runner; start explicit normalization')
        native=parse_profile(frozen['native_profile'])
        actual=processing_profile(settings,native.language,{'encoding':native.encoding,'csv_delimiter':native.csv_delimiter},probe=readiness(settings))
        expected_policy='inspection-v3-ocr' if native.version=='image-ocr-v2' else 'inspection-v2-native'
        if native!=actual or frozen['policy']!=expected_policy:
            raise ValidationError('Inspection parser/policy changed; start a new explicit normalization')
        with authorized(service.engine,org,actor,project_id,mutate=True) as (db,repo):
            package=nested_package(repo,project_id,package_id)
            if hashlib.sha256(canonical_json(package.manifest).encode()).hexdigest()!=context.payload.get('manifest_hash'):
                raise ValidationError('Package inspection input mismatch')
            identity_hash=hashlib.sha256(canonical_json({'package':package.id,'manifest':package.manifest,'parser':frozen}).encode()).hexdigest()
            existing=db.scalar(repo._query(PackageInspection).where(PackageInspection.package_id==package_id,PackageInspection.identity_hash==identity_hash))
            existing_result=existing.result if existing else None
            metadata=[pipeline._document_metadata(db,repo,project_id,package_id,doc.id) for doc in db.scalars(
                repo._query(Document).where(Document.package_id==package_id,Document.provenance_role=='document'))] if not existing else []
        if existing_result is not None:
            context.checkpoint('inspect',identity_hash,existing_result,progress=100);return
        results=[]
        from app.documents.service import DocumentService
        cancelled=DocumentService(service.engine,settings,packages=pipeline)._cancelled(context,actor)
        for doc in metadata:
            try:
                if cancelled():
                    from app.jobs.service import LeaseLost
                    raise LeaseLost()
                original=pipeline.read_original(org,project_id,doc)
                pipeline.authorize(org,actor,project_id,mutate=True)
                normalized=parser(original,doc['name'],doc['mime_type'],native,settings,cancelled)
            except StorageUnavailable:
                from app.jobs.worker import TransientInfrastructureFailure
                raise TransientInfrastructureFailure() from None
            finally:
                original=None
            unresolved=sorted({entry.state for entry in normalized.coverage if entry.state not in {'extracted','text_extracted','corrected'}})
            blockers=unresolved+['extraction_pending','tokenizer_profile_pending']
            if normalized.physical_pages is None:
                blockers.append('page_count_pending')
            results.append({'document_id':doc['id'],'page_count':normalized.physical_pages,'token_count':None,
                'token_upper_bound':None,'text_probe_bytes':sum(len((block.original_text if hasattr(block,'original_text') else block.transcription_text).encode())
                    for block in normalized.blocks if block.kind!='ocr_word'),
                'coverage':'probe' if normalized.blocks else (unresolved[0] if unresolved else 'pending'),
                'blockers':blockers,'warnings':list(normalized.warnings)})
        known_pages=sum(result['page_count'] for result in results if result['page_count'] is not None)
        pages=known_pages if all(result['page_count'] is not None for result in results) else None
        blockers=sorted({blocker for result in results for blocker in result['blockers']})
        if known_pages>settings.max_package_pages:
            blockers.append('page_limit_exceeded')
        report={'inspection_status':'inspected','analysis_allowed':False,'page_count':pages,'known_page_count':known_pages,
            'token_count':None,'token_upper_bound':None,'tokenizer_profile':'pending','parser_versions':frozen,
            'blockers':sorted(set(blockers+['extraction_pending','tokenizer_profile_pending'])),
            'warnings':sorted({warning for result in results for warning in result['warnings']}),'documents':results}
        pipeline.authorize(org,actor,project_id,mutate=True)
        def persist(writes):
            writes.add(PackageInspection(organization_id=org,project_id=project_id,package_id=package_id,
                identity_hash=identity_hash,parser_version=expected_policy,result=report))
        context.checkpoint('inspect',identity_hash,report,persist,progress=100)
    return handle


def register_handlers(registry,settings):
    from app.documents.runner import ensure_runner_available,DocumentRunnerError
    if settings.document_runner_mode=='disabled':
        return
    try:
        ensure_runner_available(settings)
    except DocumentRunnerError:
        return
    registry.register('package.inspect',inspection_handler(settings))
