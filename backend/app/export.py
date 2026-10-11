from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

STATUS_LABELS = {
    "compliant": "Соответствует", "partial": "Частично",
    "unconfirmed": "Не подтверждено", "non_compliant": "Не соответствует",
    "review": "Требует проверки",
}
IMPORTANCE_LABELS = {"mandatory": "Обязательное", "scored": "Оценочное", "desirable": "Желательное"}
COLORS = {"compliant": "E6F4EE", "partial": "FFF3CF", "unconfirmed": "FFE8D3", "non_compliant": "FDE6E6", "review": "EDF0F4"}


def safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def matrix_xlsx(check: dict) -> bytes:
    workbook = Workbook()
    matrix = workbook.active
    matrix.title = "Матрица"
    slots = sorted(check["results"], key=int)
    matrix.append(["ID", "Параметр", "Условие", "Тип"] + [safe(check["bids"][s]["name"]) for s in slots])
    by_slot = {s: {r["requirement_id"]: r for r in check["results"][s]["comparison"]["rows"]} for s in slots}
    for requirement in check["requirements"]:
        matrix.append([safe(requirement["id"]), safe(requirement["parameter"]), safe(requirement["condition"]),
                       IMPORTANCE_LABELS[requirement["importance"]]] +
                      [STATUS_LABELS[by_slot[s][requirement["id"]]["status"]] for s in slots])
        for column, slot in enumerate(slots, 5):
            matrix.cell(matrix.max_row, column).fill = PatternFill("solid", fgColor=COLORS[by_slot[slot][requirement["id"]]["status"]])

    scores = workbook.create_sheet("Оценки")
    scores.append(["Участник", "Соответствие, %", "Минимум, %", "Максимум, %", "Риск отклонения", "Обязательные требуют проверки", "Серые пункты"])
    for slot in slots:
        value = check["results"][slot]["score"]
        scores.append([safe(check["bids"][slot]["name"]), value["percent"], value["min_percent"], value["max_percent"],
                       "Да" if value["risk"] else "Нет", "Да" if value["mandatory_review"] else "Нет", value["review_count"]])

    sources = workbook.create_sheet("Источники")
    sources.append(["ID требования", "Участник / пакет", "Статус", "Обоснование", "Файл", "Место", "Цитата"])

    def append_sources(requirement_id, owner, status, explanation, citations, blocks):
        index = {b["id"]: b for b in blocks}
        for citation in citations or [None]:
            block = index.get(citation["block_id"], {}) if citation else {}
            values = [requirement_id, owner, status, explanation, block.get("filename", ""),
                      block.get("location", ""), citation["quote"] if citation else "Подтверждение не найдено"]
            sources.append([safe(value) for value in values])

    for requirement in check["requirements"]:
        append_sources(requirement["id"], "Документация заказчика", "", requirement["condition"],
                       requirement["citations"], check["blocks"])
        for slot in slots:
            row = by_slot[slot][requirement["id"]]
            append_sources(requirement["id"], check["bids"][slot]["name"], STATUS_LABELS[row["status"]],
                           row["explanation"], row["citations"], check["bids"][slot].get("blocks", []))

    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="18372D")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(65, max(16, max(len(str(c.value or "")) for c in column) + 2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
