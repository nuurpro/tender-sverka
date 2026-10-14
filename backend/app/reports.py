"""Basic version-bound checklist reports; all inputs are authorized snapshots."""
from io import BytesIO
import hashlib
import os
from pathlib import Path
from threading import Lock
from typing import Literal
from uuid import UUID
from xml.sax.saxutils import escape

from fastapi import APIRouter,Depends,Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
from sqlalchemy import select

from app.full_workspace.dependencies import scoped_repository,nested_checklist
from app.workspace.models import AuditEvent
from app.workspace.repository import Conflict,canonical_json,snapshot

router=APIRouter(prefix='/api/organizations/{organization_id}/projects/{project_id}/checklists')
IMPORTANCE={'mandatory':'Обязательное','scored':'Оценочное','desirable':'Желательное'}
_font_lock=Lock()


def source_location(citation):
    anchor=citation.get('anchor',{})
    if anchor.get('page') is not None:return 'Страница '+str(anchor['page'])
    if anchor.get('kind')=='docx':
        if anchor.get('table') is not None:return f"Таблица {anchor['table']}, строка {anchor.get('row','—')}"
        if anchor.get('paragraph') is not None:return 'Абзац '+str(anchor['paragraph'])
    if anchor.get('kind')=='xlsx':return f"Лист {anchor.get('sheet','')}, ячейка {anchor.get('cell','')}"
    if anchor.get('line_start') is not None:return 'Строка '+str(anchor['line_start'])
    return 'Блок '+str(citation.get('block_id',''))


def _safe(value):
    # Literal string cell types below preserve exact text without formula execution.
    return value


def checklist_xlsx(document):
    checklist=document['checklist'];workbook=Workbook();sheet=workbook.active;sheet.title='Требования'
    sheet.append(['ID','Область','Позиция','Категория','Параметр','Условие','Важность','Проверка'])
    sources=workbook.create_sheet('Источники');sources.append(['ID','Файл','Место','Цитата','Блок','Артефакт'])
    for row in checklist['requirements']:
        sheet.append([_safe(v) for v in [row['display_id'],'Товар' if row['scope']=='product' else 'Участник',row['item'],row['category'],row['parameter'],row['condition'],IMPORTANCE[row['importance']],'Требует проверки' if row['origin'].get('review_flags') else '']])
        for citation in row['citations']:
            sources.append([_safe(v) for v in [row['display_id'],citation.get('filename',''),source_location(citation),citation.get('quote',''),citation.get('block_id',''),citation.get('artifact_id','')]])
    metadata=workbook.create_sheet('Снимок')
    for key,value in [('Проект',document['project']['lot_name']),('Чек-лист',checklist['id']),('Версия',checklist['version']),('Редакция',checklist['revision']),('Статус',checklist['status']),('Снимок',document['snapshot_id'])]:metadata.append([key,_safe(value)])
    audit=workbook.create_sheet('Изменения');audit.append(['ID','Действие','Время','Основание'])
    for event in document['audit']:audit.append([_safe(event[key]) for key in ('id','action','created_at','reason')])
    for page in workbook:
        for line in page:
            for cell in line:
                if isinstance(cell.value,str):cell.data_type='s'
        page.freeze_panes='A2';page.auto_filter.ref=page.dimensions
        for cell in page[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245745')
        for column in page.columns:page.column_dimensions[column[0].column_letter].width=min(65,max(18,max(len(str(cell.value or '')) for cell in column)+2))
        for line in page.iter_rows(min_row=2):
            for cell in line:cell.alignment=Alignment(vertical='top',wrap_text=True)
    output=BytesIO();workbook.save(output);return output.getvalue()


def _report_font():
    with _font_lock:
        if 'TenderReport' in pdfmetrics.getRegisteredFontNames():return 'TenderReport'
        candidates=[os.environ.get('APP_REPORT_FONT_PATH',''),'/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',str(Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'/'arial.ttf')]
        path=next((Path(value) for value in candidates if value and Path(value).is_file()),None)
        if path is None:raise RuntimeError('Unicode report font is not installed; configure APP_REPORT_FONT_PATH')
        pdfmetrics.registerFont(TTFont('TenderReport',str(path)))
    return 'TenderReport'


def checklist_pdf(document):
    output=BytesIO();font=_report_font()
    body=ParagraphStyle('body',fontName=font,fontSize=10,leading=15,textColor=colors.HexColor('#24332b'),spaceAfter=7)
    heading=ParagraphStyle('heading',parent=body,fontSize=18,leading=24,spaceAfter=14)
    row_style=ParagraphStyle('requirement',parent=body,fontSize=12,leading=17,spaceBefore=13)
    note=ParagraphStyle('note',parent=body,fontSize=8,leading=12,textColor=colors.HexColor('#627067'))
    story=[]
    def paragraphs(value,style=body):
        text=str(value or '')
        for start in range(0,len(text),1500):story.append(Paragraph(escape(text[start:start+1500]).replace('\n','<br/>'),style))
    checklist=document['checklist']
    paragraphs('Чек-лист требований',heading);paragraphs(document['project']['lot_name'],row_style)
    paragraphs(f"Версия {checklist['version']} · редакция {checklist['revision']} · {'Подтверждён' if checklist['status']=='approved' else 'Черновик'}",note)
    paragraphs('Снимок: '+document['snapshot_id'],note)
    for row in checklist['requirements']:
        paragraphs(row['display_id']+' · '+str(row['parameter'] or ''),row_style)
        paragraphs(('Товар' if row['scope']=='product' else 'Участник')+' · '+IMPORTANCE[row['importance']],note)
        if row['item']:paragraphs(row['item'],note)
        paragraphs(row['condition'])
        if row['origin'].get('review_flags'):paragraphs('Требует проверки',note)
        for citation in row['citations']:
            paragraphs(str(citation.get('filename',''))+' · '+source_location(citation),note)
            paragraphs(citation.get('quote',''))
        story.append(Spacer(1,5))
    if document['audit']:
        paragraphs('Журнал изменений',row_style)
        for event in document['audit']:paragraphs(event['created_at']+' · '+event['action']+((' · '+event['reason']) if event['reason'] else ''),note)
    def page(canvas,doc):
        canvas.setFont(font,8);canvas.setFillColor(colors.HexColor('#627067'))
        canvas.drawString(42,25,'Тендер-Сверка');canvas.drawRightString(A4[0]-42,25,str(doc.page))
    SimpleDocTemplate(output,pagesize=A4,rightMargin=42,leftMargin=42,topMargin=40,bottomMargin=42,title='Чек-лист требований').build(story,onFirstPage=page,onLaterPages=page)
    return output.getvalue()


@router.get('/{checklist_id}/report.{format}')
def export_report(project_id:UUID,checklist_id:UUID,format:Literal['pdf','xlsx','json'],revision:int=Query(ge=1),repo=Depends(scoped_repository)):
    checklist=nested_checklist(repo,project_id,checklist_id)
    if checklist.revision!=revision:raise Conflict('Report revision changed; reload checklist')
    events=list(repo.session.scalars(repo._query(AuditEvent).where(AuditEvent.project_id==str(project_id),AuditEvent.entity_id==str(checklist_id)).order_by(AuditEvent.created_at,AuditEvent.id)))
    document={'project':snapshot(repo.get_project(str(project_id))),'checklist':repo._checklist_snapshot(checklist),
        'audit':[{key:snapshot(event)[key] for key in ('id','action','created_at','reason')} for event in events]}
    identity=hashlib.sha256(canonical_json(document).encode()).hexdigest();document['snapshot_id']=identity
    if format=='pdf':data=checklist_pdf(document);media='application/pdf'
    elif format=='xlsx':data=checklist_xlsx(document);media='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:data=canonical_json(document).encode();media='application/json'
    return Response(data,media_type=media,headers={'Content-Disposition':f'attachment; filename="checklist-v{checklist.version}-r{revision}.{format}"','X-Checklist-Revision':str(revision),'X-Report-Snapshot':identity})
