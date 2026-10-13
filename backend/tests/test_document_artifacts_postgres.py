"""Opt-in actual restricted PostgreSQL, private S3 and OS parser acceptance."""
import os
from uuid import uuid4

import pytest
from sqlalchemy import func,select,text
from sqlalchemy.exc import IntegrityError

from test_package_storage_postgres import actual_packages,real_storage
from app.documents.service import DocumentService,STAGE
from app.jobs.messages import envelope
from app.jobs.worker import HandlerRegistry,execute_message
from app.jobs.service import JobService
from app.packages.crypto import AuthenticationFailure
from app.packages.inputs import prepare_files
from app.workspace.database import tenant_transaction,verify_runtime_database
from app.workspace.models import (NormalizedArtifact,ArtifactUploadAttempt,Job,JobCheckpoint,Membership,Document)


@pytest.fixture
def actual_documents(actual_packages):
    packages,engine,org,actor,project=actual_packages
    mode=os.environ.get('WORKSPACE_TEST_DOCUMENT_RUNNER_MODE')
    if not mode:
        pytest.skip('Explicit actual OS document runner acceptance configuration required')
    changes={'document_runner_mode':mode}
    for key in ('document_runtime_python','document_parser_code_path','document_wsl_distribution'):
        if value:=os.environ.get('WORKSPACE_TEST_'+key.upper()):
            changes[key]=value
    settings=packages.settings.model_copy(update=changes)
    verify_runtime_database(engine)
    return DocumentService(engine,settings,packages=packages),engine,org,actor,project


def normalized(actual_documents,key='normalize',data=b'actual native artifact\n'):
    service,engine,org,actor,project=actual_documents
    original=service.packages.upload(org,actor,project,'tender',None,prepare_files([('actual.txt',data)]),'upload-'+key)['package']
    started=service.start(org,actor,project,original['id'],key,{})
    with tenant_transaction(engine,org) as db:
        job=db.get(Job,started['job']['id']);input_hash=job.input_hash
    from pydantic import SecretStr
    settings=service.settings.model_copy(update={'job_signing_secret':SecretStr('j'*48)})
    registry=HandlerRegistry('test');registry.register(STAGE,service.handle)
    message=envelope(org,started['job']['id'],input_hash,'j'*48)
    assert execute_message(engine,settings,registry,message)=='succeeded'
    assert execute_message(engine,settings,registry,message)=='duplicate'
    listed=service.artifacts(org,actor,project,original['id'],started['job']['id'])
    return original,started,listed['artifacts'][0]


def test_actual_isolated_parser_encrypted_artifact_checkpoint_rls_and_append_only(actual_documents):
    service,engine,org,actor,project=actual_documents
    original,started,artifact=normalized(actual_documents)
    result=service.blocks(org,actor,project,original['id'],artifact['id'])
    assert result['blocks'][0]['original_text']=='actual native artifact\n'
    with tenant_transaction(engine,org) as db:
        row=db.get(NormalizedArtifact,artifact['id'])
        encrypted=service.packages.store.get(row.object_key)
        assert b'actual native artifact' not in encrypted
        assert db.scalar(select(func.count()).select_from(JobCheckpoint).where(JobCheckpoint.job_id==started['job']['id']))==2
        assert db.get(ArtifactUploadAttempt,row.attempt_id).status=='committed'
    import httpx
    assert httpx.get(service.settings.s3_endpoint_url+'/'+service.settings.s3_bucket+'/'+row.object_key,timeout=10).status_code==403
    with tenant_transaction(engine,str(uuid4())) as db:
        assert not list(db.scalars(select(NormalizedArtifact)))
        assert not list(db.scalars(select(ArtifactUploadAttempt)))
    with engine.connect() as db:
        for table in ('normalized_artifacts','artifact_upload_attempts'):
            assert db.execute(text('SELECT relrowsecurity,relforcerowsecurity FROM pg_class WHERE relname=:name'),{'name':table}).one()==(True,True)
    for model,identity,field,value in [(NormalizedArtifact,row.id,'file_hash','a'*64),
        (NormalizedArtifact,row.id,'document_id',str(uuid4())),(ArtifactUploadAttempt,row.attempt_id,'status','cleaned')]:
        with pytest.raises(IntegrityError):
            with tenant_transaction(engine,org) as db:
                setattr(db.get(model,identity),field,value)
    get=service.packages.store.get
    original_cipher=get(row.object_key)
    try:
        service.packages.store.put(row.object_key,original_cipher[:-1])
        with pytest.raises(AuthenticationFailure):
            service.blocks(org,actor,project,original['id'],artifact['id'])
    finally:
        service.packages.store.put(row.object_key,original_cipher)


def test_actual_publication_ambiguous_ack_keeps_committed_artifact_and_original(actual_documents,monkeypatch):
    service,engine,org,actor,project=actual_documents
    original=service.packages.upload(org,actor,project,'tender',None,prepare_files([('a.txt',b'actual ambiguous original')]),'ambiguous-original')['package']
    response=service.start(org,actor,project,original['id'],'ambiguous-normalize',{})
    jobs=JobService(engine,org);claim=jobs.claim(response['job']['id'])
    from app.jobs.worker import ExecutionContext
    checkpoint=jobs.checkpoint
    def lost_ack(*args,**kwargs):
        result=checkpoint(*args,**kwargs)
        if args[2].startswith('doc:'):
            raise RuntimeError('test-only actual commit acknowledgment loss')
        return result
    monkeypatch.setattr(jobs,'checkpoint',lost_ack)
    service.handle(ExecutionContext(jobs,response['job']['id'],claim));jobs.finish(response['job']['id'],claim.token)
    listed=service.artifacts(org,actor,project,original['id'],response['job']['id'])
    assert len(listed['artifacts'])==1 and service.reconcile(org)==0
    assert service.blocks(org,actor,project,original['id'],listed['artifacts'][0]['id'])['blocks']
    assert service.packages.download(org,actor,project,original['id'],original['documents'][0]['id'])[0]==b'actual ambiguous original'


def test_actual_role_downgrade_during_ciphertext_put_prevents_publication_and_cleans_owned_object(actual_documents):
    service,engine,org,actor,project=actual_documents
    original=service.packages.upload(org,actor,project,'tender',None,prepare_files([('a.txt',b'actual revoke during put')]),'role-original')['package']
    response=service.start(org,actor,project,original['id'],'role-normalize',{})
    jobs=JobService(engine,org);claim=jobs.claim(response['job']['id'])
    put=service.packages.store.put
    def downgrade(key,data):
        put(key,data)
        with tenant_transaction(engine,org) as db:
            db.scalar(select(Membership).where(Membership.user_id==actor)).role='viewer'
    service.packages.store.put=downgrade
    from app.jobs.worker import ExecutionContext
    from fastapi import HTTPException
    try:
        with pytest.raises(HTTPException):
            service.handle(ExecutionContext(jobs,response['job']['id'],claim))
    finally:
        service.packages.store.put=put
    with tenant_transaction(engine,org) as db:
        assert not list(db.scalars(select(NormalizedArtifact)))
        attempt=db.scalar(select(ArtifactUploadAttempt)); assert attempt.status=='cleaned'
        original_row=db.get(Document,original['documents'][0]['id'])
    assert service.packages.store.get(original_row.object_key)
    from app.packages.storage import StorageUnavailable
    with pytest.raises(StorageUnavailable):
        service.packages.store.get(attempt.object_key)


def test_actual_retry_reparses_failed_source_retains_healthy_checkpoint_and_append_only_history(actual_documents):
    from app.documents.runner import run_document,DocumentRunnerError
    from app.workspace.repository import ValidationError
    from app.jobs.worker import ExecutionContext
    service,engine,org,actor,project=actual_documents
    original=service.packages.upload(org,actor,project,'tender',None,
        prepare_files([('a.txt',b'healthy actual source'),('b.txt',b'recovered actual source')]),'retry-original')['package']
    calls=[];failed_once=[False]
    def transient(data,name,*args,**kwargs):
        calls.append(name)
        if name=='b.txt' and not failed_once[0]:
            failed_once[0]=True
            raise DocumentRunnerError('timeout','explicit test-only transient timeout')
        return run_document(data,name,*args,**kwargs)  # Every success uses the real OS runner.
    service=DocumentService(engine,service.settings,packages=service.packages,parser=transient)
    started=service.start(org,actor,project,original['id'],'retry-stage',{})
    jobs=JobService(engine,org);claim=jobs.claim(started['job']['id'])
    with pytest.raises(ValidationError):
        service.handle(ExecutionContext(jobs,started['job']['id'],claim))
    jobs.fail(started['job']['id'],claim.token,'handler_failed')
    with tenant_transaction(engine,org) as db:
        before=list(db.scalars(select(NormalizedArtifact)))
        jobs.retry_in(db,started['job']['id'],project,actor,'actual-retry',{})
    claim=jobs.claim(started['job']['id']);service.handle(ExecutionContext(jobs,started['job']['id'],claim));jobs.finish(started['job']['id'],claim.token)
    rows=service.artifacts(org,actor,project,original['id'],started['job']['id'])['artifacts']
    assert calls==['a.txt','b.txt','b.txt']
    assert len(rows)==3 and sorted(row['attempt'] for row in rows)==[1,1,2]
    assert all(row.id in [artifact['id'] for artifact in rows] for row in before)
    with tenant_transaction(engine,org) as db:
        summaries=list(db.scalars(select(JobCheckpoint).where(JobCheckpoint.job_id==started['job']['id'],JobCheckpoint.step.like('normalized:%')).order_by(JobCheckpoint.created_at)))
        assert len(summaries)==2 and summaries[0].result!=summaries[1].result
    with pytest.raises(IntegrityError):
        with tenant_transaction(engine,org) as db:
            db.get(NormalizedArtifact,before[1].id).run_attempt=2


def test_actual_database_actor_guard_cannot_be_bypassed_by_callback_authorization_gap(actual_documents,monkeypatch):
    from app.jobs.worker import ExecutionContext
    service,engine,org,actor,project=actual_documents
    original=service.packages.upload(org,actor,project,'tender',None,prepare_files([('a.txt',b'database auth guard')]),'guard-original')['package']
    started=service.start(org,actor,project,original['id'],'guard-normal',{})
    jobs=JobService(engine,org);claim=jobs.claim(started['job']['id'])
    put=service.packages.store.put
    def downgrade(key,data):
        put(key,data)
        with tenant_transaction(engine,org) as db:
            db.scalar(select(Membership).where(Membership.user_id==actor)).role='viewer'
    monkeypatch.setattr(service.packages.store,'put',downgrade)
    monkeypatch.setattr(service.packages,'authorize',lambda *args,**kwargs: None)
    with pytest.raises(IntegrityError):
        service.handle(ExecutionContext(jobs,started['job']['id'],claim))
    with tenant_transaction(engine,org) as db:
        assert not list(db.scalars(select(NormalizedArtifact)))
        assert not list(db.scalars(select(JobCheckpoint).where(JobCheckpoint.job_id==started['job']['id'])))
        assert db.scalar(select(ArtifactUploadAttempt)).status=='cleaned'


def test_actual_sql_delay_past_lease_rolls_back_artifact_and_checkpoint(actual_documents):
    from sqlalchemy import event
    from app.jobs.worker import ExecutionContext
    from app.jobs.service import LeaseLost
    service,engine,org,actor,project=actual_documents
    original=service.packages.upload(org,actor,project,'tender',None,prepare_files([('a.txt',b'late SQL artifact')]),'lease-original')['package']
    started=service.start(org,actor,project,original['id'],'lease-normal',{})
    jobs=JobService(engine,org,lease_seconds=3);claim=jobs.claim(started['job']['id'])
    def delay(mapper,connection,target):
        connection.execute(text('SELECT pg_sleep(3.1)'))
    event.listen(NormalizedArtifact,'after_insert',delay)
    try:
        with pytest.raises(LeaseLost):
            service.handle(ExecutionContext(jobs,started['job']['id'],claim))
    finally:
        event.remove(NormalizedArtifact,'after_insert',delay)
    with tenant_transaction(engine,org) as db:
        assert not list(db.scalars(select(NormalizedArtifact)))
        assert not list(db.scalars(select(JobCheckpoint).where(JobCheckpoint.job_id==started['job']['id'])))
        assert db.scalar(select(ArtifactUploadAttempt)).status=='cleaned'


def test_actual_isolated_native_pdf_docx_sparse_formula_xlsx_and_encoded_csv(actual_documents):
    from io import BytesIO
    from docx import Document as WordDocument
    from test_document_readers import pdf_fixture,xlsx_fixture
    from app.jobs.worker import ExecutionContext
    service,engine,org,actor,project=actual_documents
    word=WordDocument();word.add_paragraph('Actual body');word.sections[0].header.paragraphs[0].text='Actual header'
    stream=BytesIO();word.save(stream)
    files=[('mixed.pdf',pdf_fixture(90,image=True)),('actual.docx',stream.getvalue()),
           ('sparse.xlsx',xlsx_fixture(cache='<v>2</v>')),('encoded.csv','имя,значение\r\nодин,два\r\n'.encode('cp1251'))]
    original=service.packages.upload(org,actor,project,'tender',None,prepare_files(files),'formats-original')['package']
    started=service.start(org,actor,project,original['id'],'formats-native',{'encoding':'cp1251'})
    jobs=JobService(engine,org);claim=jobs.claim(started['job']['id']);service.handle(ExecutionContext(jobs,started['job']['id'],claim));jobs.finish(started['job']['id'],claim.token)
    rows=service.artifacts(org,actor,project,original['id'],started['job']['id'])['artifacts']
    assert {row['format'] for row in rows}=={'pdf','docx','xlsx','csv'}
    for row in rows:
        data=service.blocks(org,actor,project,original['id'],row['id'])
        assert data['analysis_allowed'] is False and data['document_id'] in [doc['id'] for doc in original['documents']]
        if row['format']=='pdf':
            assert row['physical_pages']==1 and row['coverage_summary']['vision_pending']>0
            assert data['blocks'][0]['anchor']['rotation']==90
        elif row['format']=='docx':
            assert row['physical_pages'] is None and any('header' in block['anchor']['part'] for block in data['blocks'])
        elif row['format']=='xlsx':
            assert row['block_count']==3
            assert any('formula_cache_untrusted' in block['review_flags'] for block in data['blocks'])
        elif row['format']=='csv':
            assert data['blocks'][0]['anchor']['fields']==['имя','значение']
            assert data['blocks'][0]['anchor']['byte_start']==0
