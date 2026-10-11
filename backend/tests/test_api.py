from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook


def docx_bytes(text):
    doc = Document()
    doc.add_paragraph(text)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    from app import ai
    from app.main import create_app
    from app.schemas import ChecklistExtraction, BidComparison

    def extract(blocks):
        return ChecklistExtraction(requirements=[{
            "scope": "product", "item": None, "parameter": "RAM",
            "condition": blocks[0]["text"], "importance": "mandatory",
            "citations": [{"block_id": blocks[0]["id"], "quote": blocks[0]["text"]}],
        }])

    def compare(requirements, blocks):
        return BidComparison(rows=[{
            "requirement_id": r.id, "status": "compliant",
            "explanation": f"Check: {r.condition}",
            "citations": [{"block_id": blocks[0]["id"], "quote": blocks[0]["text"]}],
        } for r in requirements])

    monkeypatch.setattr(ai, "extract_requirements", extract)
    monkeypatch.setattr(ai, "compare_bid", compare)
    with TestClient(create_app(tmp_path)) as test_client:
        yield test_client


def create_check(client):
    response = client.post("/checks", files=[("files", ("tender.docx", docx_bytes("RAM >= 16 GB")))])
    assert response.status_code == 202
    return response.json()["id"]


def approve(client, check_id):
    data = client.get(f"/checks/{check_id}").json()
    response = client.put(f"/checks/{check_id}/requirements", json=data["requirements"])
    assert response.status_code == 200


def test_end_to_end_three_bids_source_isolation_and_export(client):
    check_id = create_check(client)
    approve(client, check_id)
    for slot in range(1, 4):
        response = client.put(f"/checks/{check_id}/bids/{slot}", data={"name": f"Bid {slot}"}, files=[("files", (f"bid{slot}.docx", docx_bytes(f"RAM {16 * slot} GB")))])
        assert response.status_code == 200
    assert client.post(f"/checks/{check_id}/compare").status_code == 202
    result = client.get(f"/checks/{check_id}").json()
    assert result["status"] == "done"
    assert len(result["results"]) == 3
    for slot in range(1, 4):
        bid_result = result["results"][str(slot)]
        assert bid_result["score"]["percent"] == 100
        assert bid_result["comparison"]["rows"][0]["citations"][0]["quote"] == f"RAM {16 * slot} GB"
    export = client.get(f"/checks/{check_id}/matrix.xlsx")
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content))
    assert "Матрица" in workbook.sheetnames
    assert "Источники" in workbook.sheetnames
    assert "Оценки" in workbook.sheetnames
    assert workbook["Оценки"]["B2"].value == 100


def test_edit_clears_results_and_is_used_in_next_comparison(client):
    check_id = create_check(client)
    approve(client, check_id)
    client.put(f"/checks/{check_id}/bids/1", data={"name": "A"}, files=[("files", ("a.docx", docx_bytes("RAM 32 GB")))])
    client.post(f"/checks/{check_id}/compare")
    reqs = client.get(f"/checks/{check_id}").json()["requirements"]
    reqs[0]["condition"] = "RAM >= 32 GB"
    assert client.put(f"/checks/{check_id}/requirements", json=reqs).status_code == 200
    assert client.get(f"/checks/{check_id}").json()["results"] == {}
    client.post(f"/checks/{check_id}/compare")
    result = client.get(f"/checks/{check_id}").json()
    assert "RAM >= 32 GB" in result["results"]["1"]["comparison"]["rows"][0]["explanation"]


def test_replacement_invalidates_matrix(client):
    check_id = create_check(client)
    approve(client, check_id)
    for _ in range(2):
        assert client.put(f"/checks/{check_id}/bids/1", data={"name": "A"}, files=[("files", ("a.docx", docx_bytes("RAM 32 GB")))]).status_code == 200
        assert client.get(f"/checks/{check_id}").json()["results"] == {}
        assert client.post(f"/checks/{check_id}/compare").status_code == 202


def test_failure_can_be_retried(client, monkeypatch):
    from app import ai
    real = ai.extract_requirements
    monkeypatch.setattr(ai, "extract_requirements", lambda blocks: (_ for _ in ()).throw(ValueError("Ошибка распознавания")))
    check_id = create_check(client)
    assert client.get(f"/checks/{check_id}").json()["status"] == "failed"
    monkeypatch.setattr(ai, "extract_requirements", real)
    assert client.post(f"/checks/{check_id}/retry").status_code == 202
    assert client.get(f"/checks/{check_id}").json()["status"] == "done"


def test_unsupported_file_is_rejected(client):
    response = client.post("/checks", files=[("files", ("scan.png", b"png"))])
    assert response.status_code == 422


def test_busy_operations_rejected(client):
    check_id = create_check(client)
    store = client.app.state.store
    assert store.claim(check_id, "compare")
    assert client.post(f"/checks/{check_id}/compare").status_code == 409
    assert client.put(f"/checks/{check_id}/requirements", json=store.get(check_id)["requirements"]).status_code == 409


def test_restart_and_atomic_claim(tmp_path):
    from app.db import Store
    store = Store(tmp_path)
    store.initialize()
    check_id = store.create([])
    with ThreadPoolExecutor(max_workers=8) as pool:
        claims = list(pool.map(lambda _: store.claim(check_id, "extract"), range(8)))
    assert sum(claims) == 1
    store.initialize()
    assert store.get(check_id)["status"] == "failed"
    assert store.get(check_id)["stage"] == "extract"


def test_export_neutralizes_formula_strings(client):
    check_id = create_check(client)
    approve(client, check_id)
    client.put(f"/checks/{check_id}/bids/1", data={"name": "=HYPERLINK(\"bad\")"}, files=[("files", ("a.docx", docx_bytes("RAM 16 GB")))])
    client.post(f"/checks/{check_id}/compare")
    workbook = load_workbook(BytesIO(client.get(f"/checks/{check_id}/matrix.xlsx").content))
    assert workbook["Матрица"]["E1"].data_type != "f"
