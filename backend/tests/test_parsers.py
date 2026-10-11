from io import BytesIO

import pytest


def test_docx_preserves_paragraph_and_table_locations(tmp_path):
    from docx import Document
    from app.parsers import parse_package
    doc = Document()
    doc.add_paragraph("RAM >= 16 GB")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Warranty"
    path = tmp_path / "offer.docx"
    doc.save(path)
    blocks = parse_package([{"id": "file1", "name": path.name, "path": str(path)}])
    assert blocks[0]["id"] == "file1:para1"
    assert any(b["id"] == "file1:t1:r1" for b in blocks)


def test_xlsx_formula_is_not_presented_as_a_fact(tmp_path):
    from openpyxl import Workbook
    from app.parsers import parse_package
    workbook = Workbook()
    workbook.active.append(["RAM", 16])
    workbook.active.append(["Budget", "=1+1"])
    path = tmp_path / "offer.xlsx"
    workbook.save(path)
    blocks = parse_package([{"id": "file1", "name": path.name, "path": str(path)}])
    assert "RAM" in blocks[0]["text"]
    assert "=1+1" not in " ".join(b["text"] for b in blocks)
    assert any("Формула не обработана" in b["text"] for b in blocks)


def test_pdf_page_number_and_empty_page_detection(tmp_path):
    from reportlab.pdfgen import canvas
    from app.parsers import parse_package
    path = tmp_path / "offer.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 700, "RAM >= 16 GB")
    pdf.showPage()
    pdf.save()
    blocks = parse_package([{"id": "file1", "name": path.name, "path": str(path)}])
    assert blocks[0]["page"] == 1
    assert blocks[0]["id"] == "file1:p1"
    blank = tmp_path / "scan.pdf"
    pdf = canvas.Canvas(str(blank))
    pdf.showPage()
    pdf.save()
    with pytest.raises(ValueError, match="текст|скан"):
        parse_package([{"id": "blank", "name": blank.name, "path": str(blank)}])


def test_package_text_limit(tmp_path):
    from docx import Document
    from app.parsers import parse_package
    doc = Document()
    doc.add_paragraph("x" * 100_001)
    path = tmp_path / "long.docx"
    doc.save(path)
    with pytest.raises(ValueError, match="100"):
        parse_package([{"id": "long", "name": path.name, "path": str(path)}])
