"""Atomic artifact publication and authorized current-source cache rebinding."""
from datetime import timedelta
from uuid import uuid4

import pytest
from sqlalchemy import select, func

from test_package_api import packages, upload
from test_full_auth_api import full_api, auth_headers
from app.workspace.database import tenant_transaction
from app.workspace.models import Job, JobCheckpoint, Membership, utcnow
from app.jobs.service import JobService, LeaseLost
from app.jobs.worker import ExecutionContext
from app.workspace.repository import Conflict, NotFound


def processing(packages):
    try:
        from app.documents.service import DocumentService
    except ImportError:
        pytest.fail('Versioned encrypted document artifact service is missing')
    from app.documents.runner import parser_code_hash
    from app.documents.readers import normalize
    client, engine, app, store, org, project = packages
    actor = app.state.auth.authenticate(client.cookies.get('tender_session')).id
    calls = []
    def parse(data, filename, mime, profile, settings, cancelled=None):
        calls.append(filename)
        return normalize(data, filename, mime, profile)
    service = DocumentService(engine, app.state.settings, packages=app.state.packages,
        parser=parse, readiness=lambda settings: {'isolated': True, 'code_hash': parser_code_hash(), 'runtime_versions': {}})
    return service, actor, calls


def package(packages, key='original', name='a.txt', data=b'alpha  beta\n'):
    client, _, _, _, org, project = packages
    return upload(client, f'/api/organizations/{org}/projects/{project}', key=key, name=name, data=data).json()['package']


def execute(service, org, response):
    jobs = JobService(service.engine, org)
    job_id = response['job']['id']
    claim = jobs.claim(job_id)
    service.handle(ExecutionContext(jobs, job_id, claim))
    jobs.finish(job_id, claim.token)
    return job_id


def test_disabled_normalization_is_explicit_and_does_not_queue(packages):
    client, engine, _, _, org, project = packages
    original = package(packages)
    path = f'/api/organizations/{org}/projects/{project}/packages/{original["id"]}/normalize'
    response = client.post(path, json={}, headers={**auth_headers(client), 'Idempotency-Key': 'normalize'})
    assert response.status_code == 503
    with tenant_transaction(engine, org) as db:
        assert db.scalar(select(func.count()).select_from(Job).where(Job.stage == 'package.normalize')) == 0


def test_publish_encrypted_payload_checkpoint_and_safe_nested_source_details(packages):
    from app.workspace import models
    service, actor, calls = processing(packages)
    _, engine, _, store, org, project = packages
    original = package(packages)
    response = service.start(org, actor, project, original['id'], 'normalize', {})
    assert service.start(org, actor, project, original['id'], 'normalize', {}) == response
    with pytest.raises(Conflict):
        service.start(org, actor, project, original['id'], 'normalize', {'encoding': 'cp1251'})
    job_id = execute(service, org, response)
    listed = service.artifacts(org, actor, project, original['id'], job_id)
    artifact = listed['artifacts'][0]
    assert artifact['document_id'] == original['documents'][0]['id']
    assert artifact['analysis_allowed'] is False and artifact['block_count'] == 1
    blocks = service.blocks(org, actor, project, original['id'], artifact['id'])
    assert blocks['blocks'][0]['original_text'] == 'alpha  beta\n'
    assert blocks['blocks'][0]['normalized_text'] == 'alpha beta'
    block = service.block(org, actor, project, original['id'], artifact['id'], blocks['blocks'][0]['id'])
    assert block['document_id'] == artifact['document_id']
    assert all(secret not in str(listed) + str(blocks) for secret in ('wrapped_key', 'object_key', 'envelope'))
    with tenant_transaction(engine, org) as db:
        row = db.get(models.NormalizedArtifact, artifact['id'])
        assert b'alpha  beta' not in store.get(row.object_key)
        attempt = db.get(models.ArtifactUploadAttempt, row.attempt_id)
        assert attempt.status == 'committed'
        assert db.scalar(select(func.count()).select_from(JobCheckpoint).where(JobCheckpoint.job_id == job_id)) == 2
    assert calls == ['a.txt']


def test_cache_rebinds_current_authorized_original_and_profile_options_invalidate(packages):
    service, actor, calls = processing(packages)
    _, _, _, _, org, project = packages
    first = package(packages)
    execute(service, org, service.start(org, actor, project, first['id'], 'first', {}))
    second = package(packages, key='second', name='renamed.txt')
    job = execute(service, org, service.start(org, actor, project, second['id'], 'second', {}))
    artifact = service.artifacts(org, actor, project, second['id'], job)['artifacts'][0]
    blocks = service.blocks(org, actor, project, second['id'], artifact['id'])
    assert blocks['document_id'] == second['documents'][0]['id']
    assert first['documents'][0]['id'] not in str(blocks)
    assert calls == ['a.txt']
    execute(service, org, service.start(org, actor, project, second['id'], 'new-profile', {'csv_delimiter': ';'}))
    assert calls == ['a.txt', 'renamed.txt']
    with pytest.raises(NotFound):
        service.blocks(org, actor, project, first['id'], artifact['id'])


def test_ambiguous_committed_checkpoint_preserves_ciphertext_and_replays(packages, monkeypatch):
    from app.workspace.models import NormalizedArtifact, ArtifactUploadAttempt
    service, actor, _ = processing(packages)
    _, engine, _, store, org, project = packages
    original = package(packages)
    response = service.start(org, actor, project, original['id'], 'ambiguous', {})
    jobs = JobService(engine, org); claim = jobs.claim(response['job']['id'])
    checkpoint = jobs.checkpoint
    def lost_ack(*args, **kwargs):
        result = checkpoint(*args, **kwargs)
        if str(args[2]).startswith('doc:'):
            raise RuntimeError('test-only lost acknowledgment')
        return result
    monkeypatch.setattr(jobs, 'checkpoint', lost_ack)
    service.handle(ExecutionContext(jobs, response['job']['id'], claim))
    jobs.finish(response['job']['id'], claim.token)
    with tenant_transaction(engine, org) as db:
        row = db.scalar(select(NormalizedArtifact))
        attempt = db.get(ArtifactUploadAttempt, row.attempt_id)
        assert attempt.status == 'committed' and store.get(row.object_key)
    assert service.reconcile(org) == 0


def test_cancel_during_storage_rolls_back_artifact_and_gc_owned_ciphertext(packages):
    from app.workspace.models import NormalizedArtifact, ArtifactUploadAttempt
    service, actor, _ = processing(packages)
    _, engine, _, store, org, project = packages
    original = package(packages)
    response = service.start(org, actor, project, original['id'], 'cancel', {})
    jobs = JobService(engine, org); claim = jobs.claim(response['job']['id'])
    def cancel():
        with tenant_transaction(engine, org) as db:
            jobs.cancel_in(db, response['job']['id'], project, actor, 'cancel', {})
    store.on_put = cancel
    with pytest.raises(LeaseLost):
        service.handle(ExecutionContext(jobs, response['job']['id'], claim))
    store.on_put = None
    with tenant_transaction(engine, org) as db:
        assert not list(db.scalars(select(NormalizedArtifact)))
        attempt = db.scalar(select(ArtifactUploadAttempt))
        assert attempt.status == 'cleaned' and attempt.object_key not in store.objects
    assert len(store.objects) == 1  # Immutable original survives.


def test_revocation_after_storage_read_blocks_artifact_payload(packages, monkeypatch):
    service, actor, _ = processing(packages)
    _, engine, _, store, org, project = packages
    original = package(packages)
    job = execute(service, org, service.start(org, actor, project, original['id'], 'read', {}))
    artifact = service.artifacts(org, actor, project, original['id'], job)['artifacts'][0]
    get = store.get
    def revoked(key):
        result = get(key)
        with tenant_transaction(engine, org) as db:
            db.delete(db.scalar(select(Membership).where(Membership.user_id == actor)))
        return result
    monkeypatch.setattr(store, 'get', revoked)
    with pytest.raises(NotFound):
        service.blocks(org, actor, project, original['id'], artifact['id'])


def test_http_generated_docx_and_xlsx_block_details_bound_to_current_original(packages):
    from io import BytesIO
    from docx import Document as WordDocument
    from openpyxl import Workbook
    from urllib.parse import quote
    service,actor,_=processing(packages)
    client,_,app,_,org,project=packages
    app.state.documents=service
    word=WordDocument();word.add_paragraph('Actual paragraph')
    book=Workbook();book.active['B7']='Actual sparse cell'
    streams=[]
    for source,name in [(word,'source.docx'),(book,'source.xlsx')]:
        stream=BytesIO();source.save(stream);streams.append((name,stream.getvalue()))
    for index,(name,data) in enumerate(streams):
        original=package(packages,key=f'http-{index}',name=name,data=data)
        base=f'/api/organizations/{org}/projects/{project}/packages/{original["id"]}'
        assert client.post(base+'/normalize',json={},headers={'Idempotency-Key':'csrf'}).status_code==403
        assert client.post(base+'/normalize',json={'actor_id':actor},headers={**auth_headers(client),'Idempotency-Key':'invalid'}).status_code==422
        started=client.post(base+'/normalize',json={},headers={**auth_headers(client),'Idempotency-Key':f'http-normal-{index}'}).json()
        execute(service,org,started)
        artifacts=client.get(base+f'/normalizations/{started["job"]["id"]}/artifacts')
        artifact=artifacts.json()['artifacts'][0]
        blocks=client.get(base+f'/artifacts/{artifact["id"]}/blocks').json()['blocks']
        detail=client.get(base+f'/artifacts/{artifact["id"]}/blocks/'+quote(blocks[0]['id'],safe=''))
        assert detail.status_code==200 and detail.headers['cache-control']=='no-store'
        assert detail.json()['document_id']==original['documents'][0]['id']
        assert detail.json()['block']['anchor']['kind'] in {'docx','xlsx'}
        assert client.get(base+f'/artifacts/{artifact["id"]}/blocks?limit=101').status_code==422
        assert client.get(base+'/normalizations').json()[0]['profile_hash']==started['profile_hash']
    with tenant_transaction(service.engine,org) as db:
        db.scalar(select(Membership).where(Membership.user_id==actor)).role='viewer'
    assert client.get(base+f'/artifacts/{artifact["id"]}').status_code==200
    assert client.post(base+'/normalize',json={},headers={**auth_headers(client),'Idempotency-Key':'viewer'}).status_code==403


def test_retry_reuses_first_document_checkpoint_and_never_cleans_its_artifact(packages):
    from app.packages.inputs import prepare_files
    from app.packages.storage import StorageUnavailable
    from app.workspace.models import NormalizedArtifact
    service,actor,calls=processing(packages)
    _,engine,_,store,org,project=packages
    original=service.packages.upload(org,actor,project,'tender',None,
        prepare_files([('a.txt',b'first'),('b.txt',b'second')]),'retry-original')['package']
    started=service.start(org,actor,project,original['id'],'retry-normal',{})
    jobs=JobService(engine,org);claim=jobs.claim(started['job']['id'])
    put=store.put
    attempt=[0]
    def second_fails(key,data):
        attempt[0]+=1
        if attempt[0]==2:
            raise StorageUnavailable()
        put(key,data)
    store.put=second_fails
    with pytest.raises(StorageUnavailable):
        service.handle(ExecutionContext(jobs,started['job']['id'],claim))
    jobs.fail(started['job']['id'],claim.token,'handler_failed');store.put=put
    with tenant_transaction(engine,org) as db:
        first=db.scalar(select(NormalizedArtifact));first_key=first.object_key
        jobs.retry_in(db,started['job']['id'],project,actor,'retry',{})
    claim=jobs.claim(started['job']['id']);service.handle(ExecutionContext(jobs,started['job']['id'],claim));jobs.finish(started['job']['id'],claim.token)
    assert calls==['a.txt','b.txt','b.txt'] and store.get(first_key)
    assert len(service.artifacts(org,actor,project,original['id'],started['job']['id'])['artifacts'])==2


def test_source_semantics_prevent_txt_csv_and_cross_project_cache_reuse(packages):
    from test_full_workspace_api import project as make_project
    service,actor,calls=processing(packages)
    client,_,_,_,org,project=packages
    txt=package(packages,key='txt',data=b'a,b\n')
    execute(service,org,service.start(org,actor,project,txt['id'],'txt',{}))
    csv=package(packages,key='csv',name='a.csv',data=b'a,b\n')
    job=execute(service,org,service.start(org,actor,project,csv['id'],'csv',{}))
    artifact=service.artifacts(org,actor,project,csv['id'],job)['artifacts'][0]
    assert service.blocks(org,actor,project,csv['id'],artifact['id'])['blocks'][0]['anchor']['fields']==['a','b']
    other=make_project(client,org)
    foreign=upload(client,f'/api/organizations/{org}/projects/{other}',key='other',name='a.txt',data=b'a,b\n').json()['package']
    execute(service,org,service.start(org,actor,other,foreign['id'],'other',{}))
    assert len(calls)==3


def test_transient_failed_parser_checkpoint_recovers_on_retry_preserving_healthy_evidence(packages):
    from app.packages.inputs import prepare_files
    from app.documents.runner import DocumentRunnerError
    from app.workspace.repository import ValidationError
    from app.workspace.models import NormalizedArtifact
    service,actor,calls=processing(packages)
    _,engine,_,store,org,project=packages
    parse=service.parser; failed_once=[False]
    def transient(data,name,*args,**kwargs):
        if name=='b.txt' and not failed_once[0]:
            failed_once[0]=True
            raise DocumentRunnerError('timeout','test-only transient parse timeout')
        return parse(data,name,*args,**kwargs)
    service.parser=transient
    original=service.packages.upload(org,actor,project,'tender',None,prepare_files([('a.txt',b'healthy'),('b.txt',b'recovered')]),'parse-original')['package']
    started=service.start(org,actor,project,original['id'],'parse-retry',{})
    jobs=JobService(engine,org);claim=jobs.claim(started['job']['id'])
    with pytest.raises(ValidationError):
        service.handle(ExecutionContext(jobs,started['job']['id'],claim))
    jobs.fail(started['job']['id'],claim.token,'handler_failed')
    with tenant_transaction(engine,org) as db:
        healthy=db.scalar(select(NormalizedArtifact).where(NormalizedArtifact.cache_eligible.is_(True)))
        healthy_id,healthy_key=healthy.id,healthy.object_key
        jobs.retry_in(db,started['job']['id'],project,actor,'retry-transient',{})
    claim=jobs.claim(started['job']['id'])
    service.handle(ExecutionContext(jobs,started['job']['id'],claim));jobs.finish(started['job']['id'],claim.token)
    artifacts=service.artifacts(org,actor,project,original['id'],started['job']['id'])['artifacts']
    assert len(artifacts)==3 and healthy_id in [row['id'] for row in artifacts]
    assert sorted(row['attempt'] for row in artifacts)==[1,1,2]
    assert calls==['a.txt','b.txt'] and store.get(healthy_key)
    assert len([row for row in artifacts if row['coverage_summary'].get('failed')])==1
