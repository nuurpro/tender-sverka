import base64
import io
import json
import re
import subprocess
import struct
import sys
import zipfile
import zlib
from pathlib import Path

import pytest


def test_trusted_unavailable_rtf_mime_cannot_become_complete_text_by_rename():
    from app.documents.readers import normalize
    from app.documents.schemas import ParserProfile
    result = normalize(br'{\rtf1\ansi original text}', 'renamed.txt', 'application/rtf', ParserProfile())
    assert result.status == 'unavailable' and not result.blocks
    assert result.coverage[0].state == 'adapter_unavailable'


def test_nested_docx_cells_count_against_actual_cell_budget():
    from docx import Document
    from io import BytesIO
    from app.documents.readers import normalize,DocumentReadError
    from app.documents.schemas import ParserProfile,ParserLimits
    doc=Document();outer=doc.add_table(rows=1,cols=1)
    outer.cell(0,0).add_table(rows=2,cols=1)
    stream=BytesIO();doc.save(stream)
    with pytest.raises(DocumentReadError,match='cell'):
        normalize(stream.getvalue(),'nested.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            ParserProfile(limits=ParserLimits(max_cells=2)))

from app.documents.readers import DocumentReadError, normalize
from app.documents.schemas import ParserLimits, ParserProfile


def package(parts):
    target = io.BytesIO()
    with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, value in parts.items():
            z.writestr(name, value)
    return target.getvalue()


def pdf_fixture(rotation=0, image=False, text=True, embedded=False):
    stream = b'BT /F1 12 Tf 100 120 Td (Native evidence) Tj ET' if text else b''
    if image:
        stream += b' q 20 0 0 20 150 180 cm /Im1 Do Q'
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
               b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               f'<< /Type /Page /Parent 2 0 R /MediaBox [20 30 320 430] /CropBox [40 50 300 410] /Rotate {rotation} /Resources << /Font << /F1 4 0 R >> /XObject << /Im1 6 0 R >> >> /Contents 5 0 R >>'.encode(),
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream',
               b'<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceGray /BitsPerComponent 8 /Length 1 >>\nstream\n\x00\nendstream']
    if embedded:
        objects[0] = b'<< /Type /Catalog /Pages 2 0 R /Names << /EmbeddedFiles << /Names [(hidden.txt) 7 0 R] >> >> >>'
        payload = b'Hidden requirement'
        objects.extend([b'<< /Type /Filespec /F (hidden.txt) /EF << /F 8 0 R >> >>',
                        b'<< /Type /EmbeddedFile /Length '+str(len(payload)).encode()+b' >>\nstream\n'+payload+b'\nendstream'])
        if embedded == 'kids':
            objects[0] = b'<< /Type /Catalog /Pages 2 0 R /Names << /EmbeddedFiles << /Kids [9 0 R] >> >> >>'
            objects.append(b'<< /Names [(hidden.txt) 7 0 R] >>')
        elif embedded in {'af', 'afpage'}:
            objects[0] = b'<< /Type /Catalog /Pages 2 0 R /AF [7 0 R] >>'
            objects[6] = b'<< /Type /Filespec /F (hidden.txt) /AFRelationship /Data /EF << /F 8 0 R >> >>'
            if embedded == 'afpage':
                objects[0] = b'<< /Type /Catalog /Pages 2 0 R >>'
                objects[2] = objects[2].replace(b'/Contents', b'/AF [7 0 R] /Contents')
    result = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f'{index} 0 obj\n'.encode()+obj+b'\nendobj\n')
    xref = len(result)
    result.extend(f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode())
    for offset in offsets[1:]:
        result.extend(f'{offset:010d} 00000 n \n'.encode())
    result.extend(f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode())
    return bytes(result)


def test_text_keeps_original_offsets_and_maps_normalized_whitespace():
    content = '  Привет\t мир\r\nВторая строка\n'.encode('utf-8')
    doc = normalize(content, 'source.txt', 'text/plain', ParserProfile())
    assert doc.physical_pages is None
    assert not doc.analysis_allowed
    assert doc.blocks[0].original_text == '  Привет\t мир\r\n'
    assert doc.blocks[0].normalized_text == 'Привет мир'
    assert doc.blocks[0].anchor.byte_start == 0
    assert doc.blocks[1].anchor.byte_start == len('  Привет\t мир\r\n'.encode())
    assert doc.blocks[0].normalization_map
    assert [b.id for b in doc.blocks] == [b.id for b in normalize(content, 'renamed.txt', 'text/plain', ParserProfile()).blocks]
    assert all(re.fullmatch(r'b-[a-f0-9]{32}', b.id) for b in doc.blocks)
    assert all(c.block_ids[0] in {b.id for b in doc.blocks} for c in doc.coverage)


def test_csv_multiline_record_has_actual_line_and_byte_span():
    source = b'name,value\r\n"two\nlines",42\r\n'
    doc = normalize(source, 'source.csv', 'text/csv', ParserProfile())
    block = doc.blocks[1]
    assert block.original_text == '"two\nlines",42\r\n'
    assert block.anchor.line_start == 2 and block.anchor.line_end == 3
    assert block.anchor.record == 2
    assert block.anchor.fields == ('two\nlines', '42')
    assert source[block.anchor.byte_start:block.anchor.byte_end].decode() == block.original_text


@pytest.mark.parametrize('separator', ['\u2028', '\u0085'])
def test_csv_unicode_separators_are_field_data_with_exact_physical_anchors(separator):
    first = f'first{separator}second,value\n'
    second = f'"quoted{separator}value\r\nnext",42\r\n'
    final = 'last,7\r'
    text = first+second+final
    source = text.encode('utf-8')
    doc = normalize(source, 'unicode.csv', 'text/csv', ParserProfile())
    assert len(doc.blocks) == 3
    assert [b.original_text for b in doc.blocks] == [first, second, final]
    assert [b.anchor.fields for b in doc.blocks] == [(f'first{separator}second', 'value'),
                                                   (f'quoted{separator}value\r\nnext', '42'), ('last', '7')]
    assert [(b.anchor.line_start, b.anchor.line_end) for b in doc.blocks] == [(1, 1), (2, 3), (4, 4)]
    for block in doc.blocks:
        assert source[block.anchor.byte_start:block.anchor.byte_end].decode('utf-8') == block.original_text
        assert text[block.anchor.char_start:block.anchor.char_end] == block.original_text


@pytest.mark.parametrize('rotation', [0, 90, 180, 270])
def test_pdf_has_original_coordinates_crop_rotation_and_mixed_coverage(rotation):
    doc = normalize(pdf_fixture(rotation, image=True), 'a.pdf', 'application/pdf', ParserProfile())
    assert doc.physical_pages == 1
    assert doc.status == 'partial' and not doc.analysis_allowed
    anchor = doc.blocks[0].anchor
    assert anchor.media_box == (20.0, 30.0, 320.0, 430.0)
    assert anchor.crop_box == (40.0, 50.0, 300.0, 410.0)
    assert anchor.rotation == rotation
    assert anchor.coordinate_space == 'pdf_original_user_space'
    x0, y0, x1, y1 = anchor.original_bbox
    assert x0 == pytest.approx(100, abs=.1)
    assert y0 == pytest.approx(117.516, abs=.1)
    assert 180 < x1 < 200 and y1 == pytest.approx(129.516, abs=.1)
    assert any(c.state == 'vision_pending' for c in doc.coverage)


def test_scan_pdf_is_pending_and_never_success():
    doc = normalize(pdf_fixture(image=True, text=False), 'scan.pdf', 'application/pdf', ParserProfile())
    assert not doc.blocks and not doc.analysis_allowed
    assert any(c.state == 'ocr_pending' for c in doc.coverage)


def test_docx_original_xml_cells_merges_and_other_parts():
    ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    doc = f'<w:document xmlns:w="{ns}"><w:body><w:p><w:r><w:t xml:space="preserve"> A  B </w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr><w:p><w:r><w:t>Merged</w:t></w:r></w:p></w:tc></w:tr><w:tr><w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr><w:p><w:r><w:t>Vertical</w:t></w:r></w:p></w:tc></w:tr><w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc></w:tr></w:tbl></w:body></w:document>'
    header = f'<w:hdr xmlns:w="{ns}"><w:p><w:r><w:t>Header evidence</w:t></w:r></w:p></w:hdr>'
    footnote = f'<w:footnotes xmlns:w="{ns}"><w:footnote w:id="1"><w:p><w:r><w:t>Footnote evidence</w:t></w:r></w:p></w:footnote></w:footnotes>'
    source = package({'word/document.xml': doc, 'word/header1.xml': header, 'word/footnotes.xml': footnote, 'word/media/image1.png': b'fake-image'})
    result = normalize(source, 'a.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', ParserProfile())
    assert result.physical_pages is None and not result.analysis_allowed
    assert result.blocks[0].original_text == ' A  B '
    assert '<' in result.blocks[0].original_xml
    merged = next(b for b in result.blocks if b.original_text == 'Merged')
    assert merged.anchor.col_span == 2
    vertical = next(b for b in result.blocks if b.original_text == 'Vertical')
    assert vertical.anchor.row_span == 2
    assert {'word/header1.xml', 'word/footnotes.xml'} <= {b.anchor.part for b in result.blocks}
    assert any(c.state == 'vision_pending' for c in result.coverage)


def docx_vertical_fixture(case='adjacent', span=1):
    def cell(text, merge=None, width=1):
        properties = f'<w:gridSpan w:val="{width}"/>' if width != 1 else ''
        if merge is not None:
            properties += '<w:vMerge/>' if merge == 'continue' else f'<w:vMerge w:val="{merge}"/>'
        return f'<w:tc><w:tcPr>{properties}</w:tcPr><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:tc>'
    first = '<w:tr>'+cell('Origin', 'restart', span)+cell('Side one')+'</w:tr>'
    if case == 'missing_column':
        second = '<w:tr><w:trPr><w:gridBefore w:val="1"/></w:trPr>'+cell('Side two')+'</w:tr>'
    elif case == 'empty_row':
        second = '<w:tr/>'
    elif case == 'changed_span':
        second = '<w:tr>'+cell('Interruption', width=2)+'</w:tr>'
    elif case == 'overlapping_restart':
        second = '<w:tr>'+cell('New origin', 'restart', 2)+'</w:tr>'
    elif case == 'ordinary_cell':
        second = '<w:tr>'+cell('Interruption')+'</w:tr>'
    elif case == 'changed_continuation':
        second = '<w:tr>'+cell('Wrong span', 'continue', 2)+'</w:tr>'
    elif case == 'shifted_continuation':
        second = '<w:tr><w:trPr><w:gridBefore w:val="1"/></w:trPr>'+cell('Wrong column', 'continue')+'</w:tr>'
    else:
        second = '<w:tr>'+cell('Continuation two', 'continue', span)+cell('Side two')+'</w:tr>'
    third = '<w:tr>'+cell('Continuation three', 'continue', span)+cell('Side three')+'</w:tr>'
    return package({'word/document.xml': '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:tbl>'+
                    first+second+third+'</w:tbl></w:body></w:document>'})


@pytest.mark.parametrize('span', [1, 2])
def test_docx_vertical_merge_requires_adjacent_matching_grid_columns(span):
    doc = normalize(docx_vertical_fixture(span=span), 'merged.docx', '', ParserProfile())
    origin = next(b for b in doc.blocks if b.original_text == 'Origin')
    assert origin.anchor.row == 1 and origin.anchor.col == 1 and origin.anchor.row_span == 3
    assert origin.anchor.col_span == span and origin.anchor.merge_origin is None
    continuation = [b for b in doc.blocks if b.original_text.startswith('Continuation')]
    assert [b.anchor.row for b in continuation] == [2, 3]
    assert all(b.anchor.merge_origin == origin.anchor.element_path and b.anchor.col_span == span for b in continuation)
    assert all(b.original_text in b.original_xml for b in [origin, *continuation])


@pytest.mark.parametrize('case', ['missing_column', 'empty_row', 'changed_span', 'overlapping_restart', 'ordinary_cell',
                                 'changed_continuation', 'shifted_continuation'])
def test_docx_vertical_merge_cannot_reconnect_across_an_interrupted_row(case):
    with pytest.raises(DocumentReadError, match='vertical merge'):
        normalize(docx_vertical_fixture(case), 'broken-merge.docx', '', ParserProfile())


def xlsx_fixture(cache=''):
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    return package({'xl/workbook.xml': f'<workbook xmlns="{ns}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Тендер" sheetId="1" r:id="rId1"/></sheets></workbook>',
                    'xl/_rels/workbook.xml.rels': '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/></Relationships>',
                    'xl/worksheets/sheet1.xml': f'<worksheet xmlns="{ns}"><dimension ref="A1:XFD1048576"/><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Merged</t></is></c><c r="B1"><f>1+1</f>{cache}</c></row><row r="1048576"><c r="XFD1048576"><v>7</v></c></row></sheetData><mergeCells><mergeCell ref="A1:A2"/></mergeCells></worksheet>'})


@pytest.mark.parametrize('cache,flag', [('', 'formula_cache_missing'), ('<v>2</v>', 'formula_cache_untrusted')])
def test_xlsx_sparse_xml_cells_raw_formula_cache_and_merge(cache, flag):
    doc = normalize(xlsx_fixture(cache), 'a.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', ParserProfile())
    assert len(doc.blocks) == 3 and doc.physical_pages is None
    assert [b.anchor.cell for b in doc.blocks] == ['A1', 'B1', 'XFD1048576']
    assert doc.blocks[0].anchor.merge_range == 'A1:A2'
    formula = doc.blocks[1]
    assert formula.anchor.raw_formula == '1+1'
    assert formula.original_text == '1+1'
    assert formula.anchor.cached_value == ('2' if cache else None)
    assert flag in formula.review_flags and doc.status == 'partial'


def test_unavailable_adapter_is_explicit():
    doc = normalize(b'CAD bytes', 'plan.dwg', 'application/acad', ParserProfile())
    assert doc.status == 'unavailable' and doc.physical_pages is None
    assert not doc.analysis_allowed and doc.coverage[0].state == 'adapter_unavailable'


def test_zip_limits_and_malformed_xml_fail_closed():
    with pytest.raises(DocumentReadError, match='ZIP'):
        normalize(package({'word/document.xml': 'x'*50000}), 'a.docx', '', ParserProfile(limits=ParserLimits(max_zip_ratio=10)))
    with pytest.raises(DocumentReadError, match='XML'):
        normalize(package({'word/document.xml': '<broken'}), 'a.docx', '', ParserProfile())
    with pytest.raises(DocumentReadError, match='cell'):
        normalize(xlsx_fixture(), 'a.xlsx', '', ParserProfile(limits=ParserLimits(max_cells=2)))


def test_profile_frozen_and_unknown_fields_rejected():
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        ParserProfile(untrusted_option=True)
    with pytest.raises(ValidationError):
        ParserProfile().language = 'en'
    with pytest.raises(TypeError):
        ParserProfile().runtime_versions['pydantic'] = 'changed'
    with pytest.raises(ValidationError):
        ParserLimits(max_cells='12')


def test_docx_utf16_exact_xml_and_unknown_top_level_content():
    ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    xml = f'<?xml version="1.0" encoding="UTF-16"?><w:document xmlns:w="{ns}"><w:body><w:p><w:r><w:t>Русский</w:t></w:r></w:p><w:altChunk/></w:body></w:document>'
    doc = normalize(package({'word/document.xml': xml.encode('utf-16')}), 'a.docx', '', ParserProfile())
    assert doc.blocks[0].original_xml == '<w:p><w:r><w:t>Русский</w:t></w:r></w:p>'
    assert any(c.state == 'adapter_unavailable' for c in doc.coverage)


def test_malformed_unhandled_xml_and_entity_payload_rejected():
    valid = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>'
    for parts in ({'word/document.xml': valid, 'word/unknown.xml': '<broken'},
                  {'word/document.xml': '<!DOCTYPE doc [<!ENTITY a "value">]><doc>&a;</doc>'}):
        with pytest.raises(DocumentReadError, match='XML'):
            normalize(package(parts), 'a.docx', '', ParserProfile())


def test_output_pixel_page_and_block_limits():
    with pytest.raises(DocumentReadError, match='output'):
        normalize(b'text', 'a.txt', '', ParserProfile(limits=ParserLimits(max_output_bytes=100)))
    with pytest.raises(DocumentReadError, match='block'):
        normalize(b'one\ntwo\n', 'a.txt', '', ParserProfile(limits=ParserLimits(max_blocks=1)))
    with pytest.raises(DocumentReadError, match='pixel'):
        normalize(pdf_fixture(image=True).replace(b'/Width 1 /Height 1', b'/Width 9999 /Height 9999'), 'a.pdf', '', ParserProfile())


def test_complete_pdf_never_claims_analysis_ready():
    result = normalize(pdf_fixture(), 'text.pdf', '', ParserProfile())
    assert result.status == 'complete' and result.physical_pages == 1
    assert not result.analysis_allowed


@pytest.mark.parametrize('embedded', [True, 'kids'])
def test_pdf_catalog_embedded_file_is_pending_without_losing_page_evidence(embedded):
    result = normalize(pdf_fixture(embedded=embedded), 'attached.pdf', 'application/pdf', ParserProfile())
    assert result.blocks[0].original_text == 'Native evidence'
    assert result.physical_pages == 1 and result.status == 'partial'
    pending = [c for c in result.coverage if c.kind == 'embedded']
    assert pending and all(c.state == 'adapter_unavailable' and c.anchor.page is None for c in pending)
    assert any('EmbeddedFiles' in c.anchor.part and 'hidden.txt' in c.anchor.part and '7' in c.anchor.part for c in pending)
    assert not result.analysis_allowed


def test_pdf_catalog_name_tree_is_bounded():
    with pytest.raises(DocumentReadError, match='structure limit'):
        normalize(pdf_fixture(embedded='kids'), 'attached.pdf', '', ParserProfile(limits=ParserLimits(max_blocks=1)))


@pytest.mark.parametrize('embedded', ['af', 'afpage'])
def test_pdf_associated_file_without_names_is_pending(embedded):
    result = normalize(pdf_fixture(embedded=embedded), 'associated.pdf', '', ParserProfile())
    assert result.blocks[0].original_text == 'Native evidence' and result.status == 'partial'
    assert any(c.kind == 'embedded' and c.state == 'adapter_unavailable' and '/AF[0]' in c.anchor.part
               and 'Filespec:7' in c.anchor.part and 'hidden.txt' in c.anchor.part for c in result.coverage)


def unknown_xlsx_fixture(location='cell'):
    with zipfile.ZipFile(io.BytesIO(xlsx_fixture())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    unknown = b'<z:evidence xmlns:z="urn:uninterpreted">Hidden requirement</z:evidence>'
    inline = b'<is><t>Known inline text</t>'+unknown+b'</is>' if location == 'inline' else b'<is><t>Known inline text</t></is>'+unknown
    if location == 'text':
        inline = b'<is><t>Known inline text'+unknown+b'</t></is>'
    if location == 'rich':
        inline = b'<is><r><rPr><b/>'+unknown+b'</rPr><t>Known inline text</t></r></is>'
    parts['xl/worksheets/sheet1.xml'] = (b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<sheetData><row r="1"><c r="A1" t="inlineStr">'+inline+b'</c></row></sheetData></worksheet>')
    if location == 'shared':
        parts['xl/worksheets/sheet1.xml'] = parts['xl/worksheets/sheet1.xml'].replace(
            b't="inlineStr">'+inline, b't="s"><v>0</v>')
        parts['xl/sharedStrings.xml'] = (b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                                       b'<si><t>Known inline text</t>'+unknown+b'</si></sst>')
    return package(parts)


@pytest.mark.parametrize('location', ['cell', 'inline', 'text', 'rich', 'shared'])
def test_xlsx_unknown_cell_xml_is_pending_and_preserves_known_text(location):
    result = normalize(unknown_xlsx_fixture(location), 'unknown.xlsx', '', ParserProfile())
    block = result.blocks[0]
    assert block.original_text == 'Known inline text'
    if location != 'shared':
        assert 'Hidden requirement' in block.original_xml
    assert 'unknown_xml_content' in block.review_flags and result.status == 'partial'
    pending = [c for c in result.coverage if c.kind == 'unknown' and c.state == 'review_pending']
    assert any(c.anchor.kind == 'xlsx' and c.anchor.cell == 'A1' and c.block_ids == (block.id,) for c in pending)
    assert not result.analysis_allowed


def unknown_xlsx_container_fixture(location):
    with zipfile.ZipFile(io.BytesIO(unknown_xlsx_fixture())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    unknown = b'<z:evidence xmlns:z="urn:uninterpreted">Hidden requirement</z:evidence>'
    sheet = parts['xl/worksheets/sheet1.xml'].replace(unknown, b'')
    if location == 'row':
        sheet = sheet.replace(b'</row>', unknown+b'</row>')
    elif location == 'row_text':
        sheet = sheet.replace(b'<row r="1">', b'<row r="1">Hidden requirement')
    elif location == 'row_tail':
        sheet = sheet.replace(b'</c>', b'</c>Hidden requirement')
    elif location == 'sheetData':
        sheet = sheet.replace(b'</sheetData>', unknown+b'</sheetData>')
    elif location == 'sheetData_text':
        sheet = sheet.replace(b'<sheetData>', b'<sheetData>Hidden requirement')
    else:
        sheet = sheet.replace(b'<sheetData>', b'<sheetViews><sheetView workbookViewId="0">'+unknown+b'</sheetView></sheetViews><sheetData>')
    parts['xl/worksheets/sheet1.xml'] = sheet
    return package(parts)


@pytest.mark.parametrize('location', ['row', 'row_text', 'row_tail', 'sheetData', 'sheetData_text', 'sheetView'])
def test_xlsx_unknown_container_xml_stays_pending_with_known_cell(location):
    result = normalize(unknown_xlsx_container_fixture(location), 'unknown-container.xlsx', '', ParserProfile())
    assert result.blocks[0].original_text == 'Known inline text' and result.status == 'partial'
    pending = [c for c in result.coverage if c.kind == 'unknown' and c.state == 'review_pending']
    assert any(c.anchor.part == 'xl/worksheets/sheet1.xml' and location.split('_')[0] in c.id for c in pending)
    if location.startswith('row'):
        assert any('row[0;r=1]' in c.id for c in pending)


def test_xlsx_known_workbook_container_does_not_hide_unknown_children():
    with zipfile.ZipFile(io.BytesIO(unknown_xlsx_container_fixture('row'))) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    parts['xl/workbook.xml'] = parts['xl/workbook.xml'].replace(
        b'</sheets>', b'<z:evidence xmlns:z="urn:uninterpreted">Hidden workbook requirement</z:evidence></sheets>')
    result = normalize(package(parts), 'unknown-workbook.xlsx', '', ParserProfile())
    assert result.blocks[0].original_text == 'Known inline text'
    assert any(c.state == 'review_pending' and c.anchor.part == 'xl/workbook.xml' and '/sheets[' in c.id
               for c in result.coverage)


def test_trusted_pdf_mime_routes_renamed_input():
    result = normalize(pdf_fixture(), 'renamed.bin', 'application/pdf', ParserProfile())
    assert result.format == 'pdf' and result.physical_pages == 1


def test_embedded_image_header_pixel_limit_in_docx_and_xlsx():
    def chunk(name, data):
        return struct.pack('>I', len(data))+name+data+struct.pack('>I', zlib.crc32(name+data))
    png = b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR', struct.pack('>IIBBBBB', 9000, 9000, 8, 0, 0, 0, 0))+chunk(b'IDAT', zlib.compress(b''))+chunk(b'IEND', b'')
    docx = {'word/document.xml': '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>', 'word/media/image.png': png}
    with zipfile.ZipFile(io.BytesIO(xlsx_fixture())) as z:
        xlsx = {name: z.read(name) for name in z.namelist()}
    xlsx['xl/media/image.png'] = png
    for format, parts in [('docx', docx), ('xlsx', xlsx)]:
        with pytest.raises(DocumentReadError, match='pixel'):
            normalize(package(parts), 'a.'+format, '', ParserProfile())


def test_actual_deflate_payload_cannot_hide_behind_forged_uncompressed_size():
    source = bytearray(package({'word/document.xml': '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>', 'word/unknown.bin': b'A'*100000}))
    cursor = 0
    while True:
        cursor = source.find(b'PK\x01\x02', cursor)
        if cursor < 0:
            break
        name_length = struct.unpack_from('<H', source, cursor+28)[0]
        name = source[cursor+46:cursor+46+name_length]
        if name == b'word/unknown.bin':
            offset = struct.unpack_from('<I', source, cursor+42)[0]
            for at in (cursor+16, offset+14):
                struct.pack_into('<I', source, at, zlib.crc32(b'AA'))
            for at in (cursor+24, offset+22):
                struct.pack_into('<I', source, at, 2)
            break
        cursor += 4
    with pytest.raises(DocumentReadError, match='ZIP'):
        normalize(bytes(source), 'a.docx', '', ParserProfile())


def test_docx_unknown_content_and_workbook_defined_names_stay_pending():
    source = package({'word/document.xml': '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:z="urn:unknown"><w:body><w:p><w:r><w:t>Known</w:t></w:r></w:p><z:evidence>Uninterpreted</z:evidence></w:body></w:document>'})
    doc = normalize(source, 'a.docx', '', ParserProfile())
    assert doc.status == 'partial' and any(c.state == 'review_pending' for c in doc.coverage)
    with zipfile.ZipFile(io.BytesIO(xlsx_fixture())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    parts['xl/workbook.xml'] = parts['xl/workbook.xml'].replace(b'</workbook>', b'<definedNames><definedName name="hidden">1+1</definedName></definedNames></workbook>')
    doc = normalize(package(parts), 'a.xlsx', '', ParserProfile())
    assert any(c.id.startswith('coverage:workbook-region:') for c in doc.coverage)


def test_real_office_writer_files_read_without_dimension_expansion():
    from docx import Document
    from openpyxl import Workbook
    word = Document()
    word.add_paragraph('Original office document')
    word.add_table(rows=1, cols=2).cell(0, 0).text = 'Cell evidence'
    output = io.BytesIO()
    word.save(output)
    parsed = normalize(output.getvalue(), 'a.docx', '', ParserProfile())
    assert any(b.original_text == 'Original office document' for b in parsed.blocks)
    workbook = Workbook()
    workbook.active['A1'] = 'Actual spreadsheet'
    workbook.active['XFD1048576'] = 7
    output = io.BytesIO()
    workbook.save(output)
    parsed = normalize(output.getvalue(), 'a.xlsx', '', ParserProfile())
    assert len(parsed.blocks) == 2
    assert parsed.blocks[0].original_text == 'Actual spreadsheet'


def test_docx_table_cell_bound_and_unknown_inline_data():
    ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    docx = f'<w:document xmlns:w="{ns}" xmlns:z="urn:unknown"><w:body><w:p><w:r><w:t>Known</w:t><z:ignored>Important</z:ignored></w:r></w:p><w:tbl><w:tr><w:tc><w:p/></w:tc><w:tc><w:p/></w:tc></w:tr></w:tbl></w:body></w:document>'
    source = package({'word/document.xml': docx})
    with pytest.raises(DocumentReadError, match='cell'):
        normalize(source, 'a.docx', '', ParserProfile(limits=ParserLimits(max_cells=1)))
    parsed = normalize(source, 'a.docx', '', ParserProfile())
    assert 'unknown_xml_content' in parsed.blocks[0].review_flags
    assert parsed.status == 'partial'


def test_entrypoint_isolated_bootstrap_and_profile_identity():
    entry = Path(__file__).resolve().parents[1] / 'app/documents/entrypoint.py'
    probe = subprocess.run([sys.executable, '-I', str(entry)], input=b'{"probe":true}', capture_output=True, check=True)
    identity = json.loads(probe.stdout)
    profile = ParserProfile(code_hash=identity['code_hash'], runtime_versions=identity['runtime_versions'])
    request = {'filename': 'a.txt', 'mime_type': 'text/plain', 'content_base64': base64.b64encode(b'actual evidence').decode(), 'profile': profile.model_dump(mode='json')}
    result = subprocess.run([sys.executable, '-I', str(entry)], input=json.dumps(request).encode(), capture_output=True, check=True)
    assert json.loads(result.stdout)['blocks'][0]['original_text'] == 'actual evidence'
    request['profile']['code_hash'] = '0'*64
    refused = subprocess.run([sys.executable, '-I', str(entry)], input=json.dumps(request).encode(), capture_output=True)
    assert refused.returncode != 0 and not refused.stdout
