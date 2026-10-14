from io import BytesIO
import json
import pytest
import pdfplumber
from openpyxl import load_workbook
from sqlalchemy import select
from test_package_api import packages
from test_analysis_service import setup,policy
from test_full_auth_api import full_api,auth_headers
from app.jobs.service import JobService
from app.jobs.worker import ExecutionContext
from app.workspace.database import tenant_transaction
from app.workspace.models import Membership


def extracted(packages):
    service,actor,original,options,_=setup(packages)
    client,engine,app,store,org,project=packages
    policy(packages,actor)
    started=service.start(org,actor,project,original['id'],'report-extract',options)
    jobs=JobService(engine,org);claim=jobs.claim(started['job']['id'])
    service.handle(ExecutionContext(jobs,started['job']['id'],claim));jobs.finish(started['job']['id'],claim.token)
    path=f'/api/organizations/{org}/projects/{project}/checklists/'+started['run']['checklist_id']
    return path,client.get(path).json(),actor


def edited(client,path,current,condition='Мощность не менее 30 Вт'):
    keys=('logical_id','display_id','scope','item','category','parameter','condition','importance','predicate','citations','origin')
    rows=[{key:row[key] for key in keys} for row in current['requirements']]
    rows[0].update(parameter='Мощность',condition=condition,predicate={'kind':'semantic','text':condition})
    response=client.put(path,json={'requirements':rows,'expected_revision':current['revision'],'reason':'Уточнение эксперта'},headers=auth_headers(client))
    assert response.status_code==200,response.text
    return response.json()


@pytest.mark.parametrize('format',['pdf','xlsx','json'])
def test_editor_saved_revision_exports_same_conditions_and_original_evidence(packages,format):
    client,*_=packages
    path,current,_=extracted(packages)
    changed=edited(client,path,current)
    response=client.get(path+'/report.'+format,params={'revision':changed['revision']})
    assert response.status_code==200,response.text
    assert response.headers['x-checklist-revision']==str(changed['revision'])
    assert len(response.headers['x-report-snapshot'])==64
    assert response.headers['cache-control']=='no-store'
    if format=='xlsx':
        workbook=load_workbook(BytesIO(response.content))
        requirements=list(workbook['Требования'].values)
        sources=list(workbook['Источники'].values)
        assert 'Мощность не менее 30 Вт' in requirements[1]
        assert 'Minimum power 25 W.' in sources[1] and 'a.txt' in sources[1]
    elif format=='pdf':
        assert response.content.startswith(b'%PDF-')
        with pdfplumber.open(BytesIO(response.content)) as pdf:
            text='\n'.join(page.extract_text() or '' for page in pdf.pages)
        assert 'Мощность не менее 30 Вт' in text
        assert 'Minimum power 25 W.' in text and 'a.txt' in text
    else:
        document=response.json()
        assert document['checklist']['requirements'][0]['condition']=='Мощность не менее 30 Вт'
        assert document['checklist']['requirements'][0]['citations']==current['requirements'][0]['citations']
        assert any(row['reason']=='Уточнение эксперта' for row in document['audit'])
    assert client.get(path+'/report.'+format,params={'revision':current['revision']}).status_code==409


def test_report_versions_remain_readable_after_approval_and_fork(packages):
    client,*_=packages
    path,current,_=extracted(packages)
    approved=client.post(path+'/approve',json={'expected_revision':current['revision']},headers=auth_headers(client)).json()
    assert approved['status']=='approved'
    fork=client.post(path+'/fork',json={},headers=auth_headers(client)).json()
    fork_path=path.rsplit('/',1)[0]+'/'+fork['id']
    edited(client,fork_path,fork)
    original=client.get(path+'/report.json',params={'revision':approved['revision']})
    assert original.status_code==200 and original.json()['checklist']['requirements'][0]['condition']=='at least 25 W'


def test_report_download_enforces_nested_scope_and_allows_read_only_viewer(packages):
    from uuid import uuid4
    client,engine,app,store,org,project=packages
    path,current,actor=extracted(packages)
    with tenant_transaction(engine,org) as db:
        db.scalar(select(Membership).where(Membership.organization_id==org,Membership.user_id==actor)).role='viewer'
    assert client.get(path+'/report.xlsx',params={'revision':current['revision']}).status_code==200
    assert client.put(path,json={'requirements':[],'expected_revision':current['revision'],'reason':'forbidden'},headers=auth_headers(client)).status_code==403
    foreign=path.replace('/projects/'+project,'/projects/'+str(uuid4()))
    assert client.get(foreign+'/report.xlsx',params={'revision':current['revision']}).status_code==404


def test_xlsx_report_treats_document_formula_prefix_as_literal_text(packages):
    client,*_=packages
    path,current,_=extracted(packages)
    changed=edited(client,path,current,condition='=HYPERLINK("https://example.invalid")')
    response=client.get(path+'/report.xlsx',params={'revision':changed['revision']})
    assert response.status_code==200
    workbook=load_workbook(BytesIO(response.content),data_only=False)
    cells=[cell for row in workbook['Требования'] for cell in row if isinstance(cell.value,str) and 'HYPERLINK' in cell.value]
    assert len(cells)==1 and cells[0].data_type!='f'
    assert cells[0].value=='=HYPERLINK("https://example.invalid")'
